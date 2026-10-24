"""Lectura de estados de cuenta bancarios de distintos bancos mexicanos.

Formatos soportados:
- CSV / TXT (lo que exporta la banca en línea de BBVA, Santander, Banorte,
  Citibanamex, HSBC, Scotiabank, Inbursa, BanBajío...): delimitador
  automático (, ; tab |), encabezados en cualquier renglón de los primeros 40.
- Excel .xlsx: primera hoja con datos.
- PDF: primero intenta leer las tablas; si no hay, recorre el texto renglón
  por renglón (modo genérico): fecha al inicio, importes al final y el
  saldo corrido para saber si cada importe fue cargo o abono.

Cada banco nombra distinto sus columnas; aquí se reconocen por sinónimos
(Fecha / Fecha operación, Concepto / Descripción, Cargo / Retiro / Débito,
Abono / Depósito / Crédito, Importe con signo, Saldo, Referencia).
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

MESES = {
    "ENE": 1, "JAN": 1, "FEB": 2, "MAR": 3, "ABR": 4, "APR": 4, "MAY": 5, "JUN": 6, "JUL": 7,
    "AGO": 8, "AUG": 8, "SEP": 9, "SET": 9, "OCT": 10, "NOV": 11, "DIC": 12, "DEC": 12,
}

BANCOS = [
    ("BBVA", ("BBVA", "BANCOMER")),
    ("Santander", ("SANTANDER",)),
    ("Banorte", ("BANORTE",)),
    ("Citibanamex", ("CITIBANAMEX", "BANAMEX")),
    ("HSBC", ("HSBC",)),
    ("Scotiabank", ("SCOTIABANK", "SCOTIA")),
    ("Inbursa", ("INBURSA",)),
    ("BanBajío", ("BANBAJIO", "BAJIO")),
    ("Afirme", ("AFIRME",)),
    ("Banregio", ("BANREGIO",)),
    ("Banco Azteca", ("AZTECA",)),
    ("Monex", ("MONEX",)),
    ("BanCoppel", ("BANCOPPEL",)),
    ("Mifel", ("MIFEL",)),
    ("Multiva", ("MULTIVA",)),
]

SINONIMOS = {
    "fecha": ("fecha operacion", "fecha de operacion", "fecha", "fecha mov", "fecha movimiento", "dia", "fecha aplicacion"),
    "descripcion": (
        "descripcion", "concepto", "detalle", "descripcion del movimiento", "movimiento", "leyenda",
        "concepto/referencia", "descripcion/concepto", "transaccion",
    ),
    "referencia": ("referencia", "ref", "no referencia", "numero de referencia", "folio", "autorizacion", "clave de rastreo", "rastreo"),
    "cargo": ("cargo", "cargos", "retiro", "retiros", "debito", "debitos", "debe", "egreso", "egresos", "salida", "salidas", "importe cargo", "cargo (-)"),
    "abono": ("abono", "abonos", "deposito", "depositos", "credito", "creditos", "haber", "ingreso", "ingresos", "entrada", "entradas", "importe abono", "abono (+)"),
    "importe": ("importe", "monto", "cantidad", "importe mxn", "monto mxn"),
    "saldo": ("saldo", "saldo final", "saldo disponible", "saldo operacion", "saldo liquidacion", "saldo mxn"),
}

PALABRAS_ABONO = (
    "DEPOSITO", "DEP ", "ABONO", "SPEI RECIBIDO", "TRANSFERENCIA RECIBIDA", "TRASPASO RECIBIDO", "RECIBIDO",
    "INTERESES A FAVOR", "INTERES A FAVOR", "DEVOLUCION", "PAGO RECIBIDO", "COBRO", "VENTAS TPV", "LIQUIDACION TPV",
)

RE_MONTO = re.compile(r"\(?-?\$?\s?\d{1,3}(?:,\d{3})*(?:\.\d{2})\)?-?|\(?-?\$?\s?\d+\.\d{2}\)?-?")
RE_FECHA_INICIO = re.compile(
    r"^\s*(\d{1,2}[/\-. ](?:\d{1,2}|[A-Za-z]{3})(?:[/\-. ](?:\d{4}|\d{2}))?|\d{4}-\d{2}-\d{2})\b"
)


@dataclass
class MovimientoLeido:
    fecha: date
    descripcion: str
    referencia: str | None
    cargo: Decimal | None
    abono: Decimal | None
    saldo: Decimal | None
    renglon: int = 0

    @property
    def tipo(self) -> str:
        return "abono" if (self.abono or Decimal("0")) > 0 else "cargo"

    @property
    def monto(self) -> Decimal:
        return (self.abono if self.tipo == "abono" else self.cargo) or Decimal("0")


@dataclass
class ResultadoLectura:
    movimientos: list[MovimientoLeido]
    formato: str
    saldo_inicial: Decimal | None = None
    saldo_final: Decimal | None = None
    banco_detectado: str | None = None
    advertencias: list[str] = field(default_factory=list)


class ErrorLectura(ValueError):
    """El archivo no se pudo interpretar como estado de cuenta."""


# --------------------------------------------------------------------------
# Utilidades de normalización
# --------------------------------------------------------------------------

def _norm(texto) -> str:
    if texto is None:
        return ""
    texto = str(texto).strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"[\.\:\#\$\*]", "", texto)
    return re.sub(r"\s+", " ", texto).strip()


def parse_monto(valor) -> Decimal | None:
    """'$1,234.56' '(1,234.56)' '-1234.56' '1234.56-' '1.234,56' -> Decimal."""
    if valor is None:
        return None
    if isinstance(valor, (int, float, Decimal)):
        try:
            return Decimal(str(valor)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None
    texto = str(valor).strip().upper().replace("MXN", "").replace("MN", "").replace("$", "").replace(" ", "")
    if not texto or texto in ("-", "--"):
        return None
    negativo = False
    if texto.startswith("(") and texto.endswith(")"):
        negativo, texto = True, texto[1:-1]
    if texto.endswith("-"):
        negativo, texto = True, texto[:-1]
    if texto.startswith("-"):
        negativo, texto = True, texto[1:]
    if texto.endswith("CR"):
        texto = texto[:-2]
    if texto.endswith("DR") or texto.endswith("DB"):
        negativo, texto = True, texto[:-2]
    # Formato europeo 1.234,56
    if re.fullmatch(r"\d{1,3}(\.\d{3})+,\d{2}", texto) or re.fullmatch(r"\d+,\d{2}", texto):
        texto = texto.replace(".", "").replace(",", ".")
    texto = texto.replace(",", "")
    if not re.fullmatch(r"\d+(\.\d+)?", texto):
        return None
    try:
        numero = Decimal(texto).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None
    return -numero if negativo else numero


def parse_fecha(valor, anio_defecto: int | None = None, mes_defecto: int | None = None) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and 20000 < float(valor) < 80000:
        # Número de serie de Excel
        return date(1899, 12, 30) + timedelta(days=int(valor))
    texto = str(valor).strip().upper()
    texto = re.sub(r"\s+", " ", texto)
    texto = texto.split(" ")[0] if re.match(r"^\d{4}-\d{2}-\d{2}[ T]", texto) else texto

    m = re.fullmatch(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", texto)
    if m:
        return _fecha_segura(int(m.group(1)), int(m.group(2)), int(m.group(3)))

    m = re.fullmatch(r"(\d{1,2})[/\-. ]([A-Z]{3})[A-Z]*\.?(?:[/\-. ](\d{2,4}))?", texto)
    if m:
        mes = MESES.get(m.group(2)[:3])
        if not mes:
            return None
        anio = _anio(m.group(3), anio_defecto, mes, mes_defecto)
        return _fecha_segura(anio, mes, int(m.group(1))) if anio else None

    m = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})(?:[/\-.](\d{2,4}))?", texto)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        dia, mes = (a, b)
        if b > 12 and a <= 12:  # vino como mm/dd
            dia, mes = b, a
        anio = _anio(m.group(3), anio_defecto, mes, mes_defecto)
        return _fecha_segura(anio, mes, dia) if anio else None
    return None


def _anio(texto: str | None, anio_defecto: int | None, mes: int, mes_defecto: int | None) -> int | None:
    if texto:
        anio = int(texto)
        return anio + 2000 if anio < 100 else anio
    if anio_defecto is None:
        return None
    # Estado de cuenta de enero con movimientos del 31 de diciembre anterior.
    if mes_defecto and mes == 12 and mes_defecto == 1:
        return anio_defecto - 1
    return anio_defecto


def _fecha_segura(anio: int, mes: int, dia: int) -> date | None:
    try:
        return date(anio, mes, dia)
    except ValueError:
        return None


def detectar_banco(texto: str) -> str | None:
    t = texto.upper()
    for nombre, claves in BANCOS:
        if any(c in t for c in claves):
            return nombre
    return None


# --------------------------------------------------------------------------
# Interpretación de una tabla (CSV, Excel o tabla de PDF)
# --------------------------------------------------------------------------

def _columna_de(celda: str) -> str | None:
    n = _norm(celda)
    if not n:
        return None
    for campo in ("cargo", "abono", "saldo", "referencia", "fecha", "descripcion", "importe"):
        for sinonimo in SINONIMOS[campo]:
            if n == sinonimo or n.startswith(sinonimo + " ") or n.endswith(" " + sinonimo):
                return campo
    return None


def _encontrar_encabezado(filas: list[list]) -> tuple[int, dict[str, int]] | None:
    for idx, fila in enumerate(filas[:40]):
        mapa: dict[str, int] = {}
        for col, celda in enumerate(fila):
            campo = _columna_de(celda)
            if campo and campo not in mapa:
                mapa[campo] = col
        tiene_monto = any(k in mapa for k in ("cargo", "abono", "importe"))
        if "fecha" in mapa and tiene_monto and len(mapa) >= 3:
            return idx, mapa
    return None


def interpretar_tabla(filas: list[list], periodo: str | None) -> list[MovimientoLeido]:
    anio, mes = _periodo(periodo)
    encabezado = _encontrar_encabezado(filas)
    if encabezado is None:
        raise ErrorLectura(
            "No encontramos los encabezados del estado de cuenta. El archivo debe tener columnas como "
            "Fecha, Concepto/Descripción y Cargo/Abono (o Importe) -- así lo exporta la banca en línea."
        )
    idx_encabezado, mapa = encabezado
    movimientos: list[MovimientoLeido] = []
    for num, fila in enumerate(filas[idx_encabezado + 1 :], start=idx_encabezado + 2):
        def celda(campo):
            col = mapa.get(campo)
            return fila[col] if col is not None and col < len(fila) else None

        fecha = parse_fecha(celda("fecha"), anio, mes)
        if fecha is None:
            # Renglón de continuación (descripción en dos líneas) o de totales.
            if movimientos and celda("descripcion") and not any(
                parse_monto(celda(c)) for c in ("cargo", "abono", "importe")
            ):
                movimientos[-1].descripcion = f"{movimientos[-1].descripcion} {str(celda('descripcion')).strip()}".strip()
            continue
        cargo = parse_monto(celda("cargo"))
        abono = parse_monto(celda("abono"))
        if cargo is None and abono is None and "importe" in mapa:
            importe = parse_monto(celda("importe"))
            if importe is not None:
                if importe < 0:
                    cargo = -importe
                else:
                    abono = importe
        # Algunos bancos ponen el cargo con signo negativo en su columna.
        if cargo is not None and cargo < 0:
            cargo = -cargo
        if abono is not None and abono < 0:
            cargo, abono = (cargo or Decimal("0")) + (-abono), None
        if not cargo and not abono:
            continue
        descripcion = " ".join(str(celda("descripcion") or "").split())
        referencia = celda("referencia")
        movimientos.append(
            MovimientoLeido(
                fecha=fecha,
                descripcion=descripcion,
                referencia=(" ".join(str(referencia).split())[:120] or None) if referencia not in (None, "") else None,
                cargo=cargo if cargo else None,
                abono=abono if abono else None,
                saldo=parse_monto(celda("saldo")),
                renglon=num,
            )
        )
    if not movimientos:
        raise ErrorLectura("El archivo tiene encabezados pero ningún movimiento con fecha e importe.")
    return movimientos


def _periodo(periodo: str | None) -> tuple[int | None, int | None]:
    if periodo and re.fullmatch(r"\d{4}-\d{2}", periodo):
        return int(periodo[:4]), int(periodo[5:])
    return None, None


# --------------------------------------------------------------------------
# Lectores por formato
# --------------------------------------------------------------------------

def _decodificar(contenido: bytes) -> str:
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return contenido.decode("utf-8", errors="replace")


def leer_csv(contenido: bytes, periodo: str | None) -> ResultadoLectura:
    texto = _decodificar(contenido)
    muestra = texto[:5000]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
        delimitador = dialecto.delimiter
    except csv.Error:
        delimitador = max(",;\t|", key=muestra.count)
    filas = [fila for fila in csv.reader(io.StringIO(texto), delimiter=delimitador)]
    movimientos = interpretar_tabla(filas, periodo)
    return ResultadoLectura(movimientos=movimientos, formato="csv", banco_detectado=detectar_banco(muestra))


def leer_xlsx(contenido: bytes, periodo: str | None) -> ResultadoLectura:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ErrorLectura("Falta la librería openpyxl en el servidor para leer Excel.") from exc
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ErrorLectura(
            "No se pudo abrir el Excel. Si es un .xls antiguo, ábrelo y guárdalo como .xlsx o .csv."
        ) from exc
    ultimo_error: ErrorLectura | None = None
    for hoja in libro.worksheets:
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
        try:
            movimientos = interpretar_tabla(filas, periodo)
        except ErrorLectura as exc:
            ultimo_error = exc
            continue
        texto = " ".join(str(c) for f in filas[:15] for c in f if c)
        return ResultadoLectura(movimientos=movimientos, formato="xlsx", banco_detectado=detectar_banco(texto))
    raise ultimo_error or ErrorLectura("El Excel no tiene hojas con movimientos.")


def leer_pdf(contenido: bytes, periodo: str | None) -> ResultadoLectura:
    import pdfplumber

    advertencias: list[str] = []
    try:
        pdf = pdfplumber.open(io.BytesIO(contenido))
    except Exception as exc:  # noqa: BLE001
        raise ErrorLectura("No se pudo abrir el PDF (¿está protegido con contraseña?).") from exc

    with pdf:
        textos = [(pagina.extract_text() or "") for pagina in pdf.pages]
        texto_completo = "\n".join(textos)
        banco = detectar_banco(texto_completo[:3000])

        # 1) Tablas
        filas_tablas: list[list] = []
        for pagina in pdf.pages:
            for tabla in pagina.extract_tables() or []:
                filas_tablas.extend([[(c or "").replace("\n", " ") for c in fila] for fila in tabla])
        if filas_tablas:
            try:
                movimientos = interpretar_tabla(filas_tablas, periodo)
                saldo_inicial = _buscar_saldo(texto_completo, ("SALDO ANTERIOR", "SALDO INICIAL"))
                return ResultadoLectura(
                    movimientos=movimientos,
                    formato="pdf",
                    banco_detectado=banco,
                    saldo_inicial=saldo_inicial,
                    advertencias=["PDF leído desde sus tablas: revisa que cargos y abonos estén bien."],
                )
            except ErrorLectura:
                pass

    # 2) Texto renglón por renglón
    if not texto_completo.strip():
        raise ErrorLectura("El PDF no tiene texto seleccionable (parece escaneado). Descarga el CSV/Excel del banco.")
    movimientos, saldo_inicial = _leer_texto_pdf(texto_completo, periodo)
    if not movimientos:
        raise ErrorLectura(
            "No reconocimos movimientos en el PDF. Descarga el estado de cuenta en CSV o Excel desde la banca en línea."
        )
    advertencias.append(
        "PDF leído en modo genérico (beta): se infirió cargo/abono con el saldo corrido. Revisa los movimientos."
    )
    return ResultadoLectura(
        movimientos=movimientos,
        formato="pdf",
        banco_detectado=banco,
        saldo_inicial=saldo_inicial,
        advertencias=advertencias,
    )


def _buscar_saldo(texto: str, etiquetas: tuple[str, ...]) -> Decimal | None:
    for linea in texto.upper().splitlines():
        if any(e in linea for e in etiquetas):
            montos = RE_MONTO.findall(linea)
            if montos:
                return parse_monto(montos[-1])
    return None


def _es_abono_por_texto(descripcion: str) -> bool:
    d = f" {descripcion.upper()} "
    return any(p in d for p in PALABRAS_ABONO)


def _leer_texto_pdf(texto: str, periodo: str | None) -> tuple[list[MovimientoLeido], Decimal | None]:
    anio, mes = _periodo(periodo)
    saldo_inicial = _buscar_saldo(texto, ("SALDO ANTERIOR", "SALDO INICIAL"))
    saldo_corrido = saldo_inicial
    movimientos: list[MovimientoLeido] = []
    for num, linea in enumerate(texto.splitlines(), start=1):
        linea = linea.strip()
        if not linea:
            continue
        m = RE_FECHA_INICIO.match(linea)
        fecha = parse_fecha(m.group(1), anio, mes) if m else None
        if fecha is None:
            if movimientos and not RE_MONTO.search(linea) and len(linea) < 120:
                movimientos[-1].descripcion = f"{movimientos[-1].descripcion} {linea}".strip()
            continue
        resto = linea[m.end():].strip()
        # BBVA repite la fecha (operación y liquidación): se descarta la segunda.
        m2 = RE_FECHA_INICIO.match(resto)
        if m2 and parse_fecha(m2.group(1), anio, mes):
            resto = resto[m2.end():].strip()
        montos_txt = RE_MONTO.findall(resto)
        montos = [parse_monto(x) for x in montos_txt]
        montos = [x for x in montos if x is not None]
        if not montos:
            continue
        descripcion = resto
        for x in montos_txt:
            descripcion = descripcion.replace(x, " ")
        descripcion = " ".join(descripcion.split())
        if "SALDO" in descripcion.upper() and len(montos) == 1 and "ANTERIOR" in descripcion.upper():
            continue

        saldo = None
        if len(montos) >= 2:
            monto, saldo = abs(montos[-2]), montos[-1]
        else:
            monto = abs(montos[0])
        es_abono: bool
        if saldo is not None and saldo_corrido is not None:
            if abs(saldo_corrido + monto - saldo) <= Decimal("0.02"):
                es_abono = True
            elif abs(saldo_corrido - monto - saldo) <= Decimal("0.02"):
                es_abono = False
            else:
                es_abono = _es_abono_por_texto(descripcion)
        else:
            es_abono = _es_abono_por_texto(descripcion)
        if saldo is not None:
            saldo_corrido = saldo
        elif saldo_corrido is not None:
            saldo_corrido = saldo_corrido + monto if es_abono else saldo_corrido - monto
        movimientos.append(
            MovimientoLeido(
                fecha=fecha,
                descripcion=descripcion,
                referencia=None,
                cargo=None if es_abono else monto,
                abono=monto if es_abono else None,
                saldo=saldo,
                renglon=num,
            )
        )
    return movimientos, saldo_inicial


def leer_estado_cuenta(contenido: bytes, nombre_archivo: str, periodo: str | None) -> ResultadoLectura:
    nombre = (nombre_archivo or "").lower()
    if nombre.endswith(".pdf") or contenido[:4] == b"%PDF":
        resultado = leer_pdf(contenido, periodo)
    elif nombre.endswith((".xlsx", ".xlsm")) or contenido[:2] == b"PK":
        resultado = leer_xlsx(contenido, periodo)
    elif nombre.endswith(".xls"):
        raise ErrorLectura("Los .xls antiguos no se pueden leer: guárdalo como .xlsx o .csv desde Excel.")
    else:
        resultado = leer_csv(contenido, periodo)

    resultado.banco_detectado = resultado.banco_detectado or detectar_banco(nombre_archivo or "")
    movs = resultado.movimientos
    if movs and resultado.saldo_inicial is None and movs[0].saldo is not None:
        # Saldo antes del primer movimiento = su saldo deshaciendo su importe.
        primero = movs[0]
        resultado.saldo_inicial = primero.saldo - primero.monto if primero.tipo == "abono" else primero.saldo + primero.monto
    if movs and resultado.saldo_final is None:
        con_saldo = [m for m in movs if m.saldo is not None]
        if con_saldo:
            resultado.saldo_final = con_saldo[-1].saldo
    # Movimientos fuera del mes del estado de cuenta: se avisan, no se descartan.
    anio, mes = _periodo(periodo)
    if anio and mes:
        fuera = [m for m in movs if (m.fecha.year, m.fecha.month) != (anio, mes)]
        if fuera:
            resultado.advertencias.append(
                f"{len(fuera)} movimiento(s) tienen fecha fuera de {periodo}; revisa que el periodo elegido sea el correcto."
            )
    return resultado
